"""Official source extraction. A fetched page alone is never a verified fact."""
from __future__ import annotations

import asyncio
import io
import json
import re
import unicodedata
from datetime import datetime, timezone
from urllib.parse import urljoin, urlparse

import httpx
from bs4 import BeautifulSoup
from cachetools import TTLCache
from openpyxl import load_workbook

HOME_AFFAIRS = "https://immi.homeaffairs.gov.au/visas/working-in-australia/skillselect/invitation-rounds"
JSA_SHORTAGE = "https://www.jobsandskills.gov.au/data/occupation-shortage"
STATES = {"ACT": "Australian Capital Territory", "NSW": "New South Wales", "NT": "Northern Territory", "QLD": "Queensland", "SA": "South Australia", "TAS": "Tasmania", "VIC": "Victoria", "WA": "Western Australia"}
_source_cache = TTLCache(maxsize=4, ttl=3600)
_locks: dict[str, asyncio.Lock] = {}


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def clean(value):
    return " ".join("".join(c for c in str(value or "") if unicodedata.category(c) != "Cf").split())


def norm(value):
    return re.sub(r"[^a-z0-9]+", " ", clean(value).lower()).strip()


def state_code(value):
    value = norm(value)
    return next((code for code, name in STATES.items() if value in {code.lower(), name.lower()}), None)


def source_meta(name, url):
    return {"source": name, "sourceUrl": url, "checkedAt": now_iso(), "updated": "just now"}


async def fetch(url, binary=False):
    async with httpx.AsyncClient(timeout=httpx.Timeout(35, connect=10), follow_redirects=True,
                                 headers={"User-Agent": "Pathway/0.2 (public data research)"}) as client:
        response = await client.get(url)
        response.raise_for_status()
        return response.content if binary else response.text


def expand_home_affairs(html):
    """Home Affairs stores its public accordion HTML in PageSchemaHiddenField JSON."""
    soup = BeautifulSoup(html, "html.parser")
    blocks = []
    for field in soup.select('input[type="hidden"][value]'):
        if "PageSchemaHiddenField" not in (field.get("id", "") + field.get("name", "")):
            continue
        try:
            schema = json.loads(field["value"])
        except (ValueError, TypeError):
            continue
        for item in schema.get("content", []) if isinstance(schema, dict) else []:
            if isinstance(item, dict) and isinstance(item.get("block"), str):
                blocks.append(item["block"])
    return BeautifulSoup(str(soup) + "\n".join(blocks), "html.parser")


def integer(value):
    value = clean(value).replace(",", "")
    return int(value) if re.fullmatch(r"\d+", value) else None


def parse_round(html):
    soup = expand_home_affairs(html)
    text = clean(soup.get_text(" ", strip=True))
    dated = []
    for heading in soup.find_all(re.compile(r"^h[1-6]$")):
        match = re.search(r"Invitations issued on\s+(\d{1,2}\s+[A-Za-z]+\s+20\d{2})", clean(heading.get_text(" ")), re.I)
        if match:
            try:
                dated.append((datetime.strptime(match[1], "%d %B %Y"), heading, match[1]))
            except ValueError:
                pass
    if not dated:
        raise ValueError("round_date_missing")
    _, heading, date = max(dated, key=lambda entry: entry[0])
    totals = None
    tie_break = None
    rows = []
    # Only the selected round section: never read a previous round's score as current.
    for element in heading.next_elements:
        if getattr(element, "name", None) and re.fullmatch(r"h[1-6]", element.name):
            heading_text = clean(element.get_text(" "))
            if re.search(r"Invitations issued on|Total invitations issued|State and Territory", heading_text, re.I):
                break
        if getattr(element, "name", None) != "table":
            continue
        table_rows = [[clean(c.get_text(" ", strip=True)) for c in row.find_all(["td", "th"], recursive=False)] for row in element.find_all("tr")]
        if not table_rows:
            continue
        header = " ".join(table_rows[0]).lower()
        if "total" in header and ("invited" in header or "invitations" in header):
            for row in table_rows[1:]:
                if len(row) >= 2 and "189" in row[0]:
                    totals = integer(row[1])
                    tie_break = row[2] if len(row) > 2 else None
        if "occupation" in header:
            # Legacy tables have separate 189 and 491 columns. Select 189 explicitly.
            score_column = next((i for i, cell in enumerate(table_rows[0]) if "189" in cell), None)
            for row in table_rows[1:]:
                col = score_column if score_column is not None else (1 if len(row) == 2 else None)
                if col is not None and len(row) > col:
                    rows.append({"occupation": row[0].rstrip("* "), "minimumPoints": integer(row[col]), "publishedValue": row[col]})
    if totals is None:
        raise ValueError("round_invitation_count_missing")
    next_round = re.search(r"next invitation round[^.]{0,250}?(?:held by|held on|by)\s+(\d{1,2}\s+[A-Za-z]+\s+20\d{2})", text, re.I)
    return {"latestRound": {"date": date, "invitations": totals, "tieBreak": tie_break},
            "occupationRows": rows, "nextRound": next_round[1] if next_round else None,
            "occupationTableParsed": bool(rows)}


def rating(value):
    key = norm(value)
    return {"s": "Shortage", "shortage": "Shortage", "ns": "No shortage", "no shortage": "No shortage",
            "r": "Regional shortage", "regional shortage": "Regional shortage",
            "m": "Metropolitan shortage", "metropolitan shortage": "Metropolitan shortage"}.get(key)


def parse_osl(raw):
    """Keep code namespaces separate and decode actual S / NS / R / M ratings."""
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    records = []
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            header = None
            for _ in range(40):
                row = next(rows, None)
                if row is None:
                    break
                values = [norm(v) for v in row]
                if any("national" in v or v == "australia" for v in values) and any("occupation" in v or "title" in v for v in values):
                    header = values
                    break
            if header is None:
                continue
            national = next((i for i, h in enumerate(header) if "national" in h or h == "australia"), None)
            state_columns = {code: next((i for i, h in enumerate(header) if h in {code.lower(), name.lower(), code.lower() + " rating", name.lower() + " rating"}), None) for code, name in STATES.items()}
            for row in rows:
                cell = lambda i: clean(row[i]) if i is not None and i < len(row) else ""
                for namespace in ("ANZSCO", "OSCA"):
                    columns = [i for i, h in enumerate(header) if namespace.lower() in h]
                    # Some workbooks label columns simply Code / Occupation on named sheets.
                    if not columns and namespace.lower() in sheet.title.lower():
                        columns = [i for i, h in enumerate(header) if "code" in h or "occupation" in h]
                    code_col = next((i for i in columns if "code" in header[i] or header[i] == namespace.lower() or ("6 digit" in header[i] and "title" not in header[i] and "occupation" not in header[i])), None)
                    name_col = next((i for i in columns if "code" not in header[i] and ("occupation" in header[i] or "title" in header[i])), None)
                    if name_col is None and code_col is not None:
                        name_col = next((i for i, h in enumerate(header) if h in {"occupation", "occupation title"}), None)
                    code = cell(code_col).removesuffix(".0")
                    if not re.fullmatch(r"\d{6}", code) or not cell(name_col):
                        continue
                    records.append({"occupation": cell(name_col), "code": code, "classification": namespace,
                                    "nationalRating": rating(cell(national)),
                                    "stateRatings": {s: rating(cell(i)) for s, i in state_columns.items()}})
    finally:
        workbook.close()
    if not records:
        raise ValueError("osl_columns_not_recognised")
    return records


def match_osl(records, occupation, state, anzsco="", osca=""):
    # ANZSCO is the migration namespace; never use an OSCA code as ANZSCO.
    for namespace, code in (("ANZSCO", anzsco), ("OSCA", osca)):
        if code:
            matches = [r for r in records if r["classification"] == namespace and r["code"] == clean(code)]
            if matches:
                result = dict(matches[0], matchType="code")
                break
    else:
        matches = [r for r in records if norm(r["occupation"]) == norm(occupation)]
        matches = [r for r in matches if r["classification"] == "ANZSCO"] or matches
        if len({(r["classification"], r["code"]) for r in matches}) != 1:
            return None
        result = dict(matches[0], matchType="exact_title")
    code = state_code(state)
    return dict(result, state=code or state, stateRating=result["stateRatings"].get(code))


def discover_osl(html):
    soup = BeautifulSoup(html, "html.parser")
    links = []
    for link in soup.find_all("a", href=True):
        url = urljoin(JSA_SHORTAGE, link["href"])
        label = norm(url + " " + link.get_text(" "))
        if urlparse(url).hostname != "www.jobsandskills.gov.au":
            continue
        if not urlparse(url).path.lower().endswith(".xlsx") or not all(word in label for word in ("occupation", "shortage", "list")):
            continue
        if "4 digit" in label or "unit group" in label:
            continue
        years = re.findall(r"20\d{2}", url + " " + link.get_text(" "))
        if years:
            links.append((max(map(int, years)), url))
    if not links:
        raise ValueError("osl_download_missing")
    return max(links, key=lambda pair: pair[0])


async def official_dataset(kind):
    if kind in _source_cache:
        return _source_cache[kind]
    lock = _locks.setdefault(kind, asyncio.Lock())
    async with lock:
        if kind in _source_cache:
            return _source_cache[kind]
        if kind == "round":
            result = parse_round(await fetch(HOME_AFFAIRS))
        else:
            year, url = discover_osl(await fetch(JSA_SHORTAGE))
            result = {"records": parse_osl(await fetch(url, binary=True)), "oslYear": year, "downloadUrl": url}
        result["checkedAt"] = now_iso()
        _source_cache[kind] = result
        return result


async def shortage_intelligence(occupation, state, anzsco="", osca=""):
    meta = source_meta("Jobs and Skills Australia", JSA_SHORTAGE)
    try:
        dataset = await official_dataset("osl")
        match = match_osl(dataset["records"], occupation, state, anzsco, osca)
        known = bool(match and match["nationalRating"])
        return {**meta, "checkedAt": dataset["checkedAt"], "status": "fresh" if known else "partial",
                "freshness": "Published workbook loaded", "oslYear": dataset["oslYear"], "downloadUrl": dataset["downloadUrl"],
                "occupationChecked": occupation, "occupationResult": match,
                "occupationMatch": match["nationalRating"] in {"Shortage", "Regional shortage", "Metropolitan shortage"} if known else None,
                "shortage": match["nationalRating"] if known else "Not matched",
                "shortageNote": (f"Published rating for {match['occupation']} ({match['classification']} {match['code']}). The selected classification has not been assessed against your duties."
                                 if known else "No exact occupation title or supplied classification code was matched in the published workbook. Confirm your occupation code.")}
    except Exception as exc:
        return {**meta, "status": "unavailable", "freshness": "No verified rating", "shortage": "Unavailable", "occupationMatch": None,
                "shortageNote": "The official shortage workbook could not be retrieved or its columns could not be read.", "error": type(exc).__name__}


async def migration_intelligence(occupation, state, anzsco=""):
    meta = source_meta("Department of Home Affairs", HOME_AFFAIRS)
    try:
        dataset = await official_dataset("round")
    except Exception as exc:
        return {**meta, "status": "unavailable", "freshness": "No verified round", "latestRound": {},
                "note": "The official round tables could not be retrieved or parsed. Open Home Affairs to check the published result.", "error": type(exc).__name__}
    target = occupation
    mapping = None
    rows = dataset["occupationRows"]
    match = next((row for row in rows if norm(row["occupation"]) == norm(target)), None)
    mapping_note = None
    # A different free text role must not silently receive a neighbouring occupation's score.
    if not match and anzsco:
        try:
            osl = await official_dataset("osl")
            mapping = match_osl(osl["records"], "", state, anzsco=anzsco)
            if mapping:
                target = mapping["occupation"]
                match = next((row for row in rows if norm(row["occupation"]) == norm(target)), None)
        except Exception:
            mapping_note = "The supplied ANZSCO code could not be resolved to an official title."
    points = match["minimumPoints"] if match else None
    if points is not None:
        note = f"Published minimum for {match['occupation']} in this historical round. This is not your points total or a forecast."
    elif not dataset["occupationTableParsed"]:
        note = "The round total was extracted, but the occupation score table could not be read."
    elif match:
        note = f"Home Affairs lists {match['occupation']} with {match['publishedValue'] or 'no numeric minimum'} in this round."
    else:
        note = f"No published score row matched {target} in this round. This does not establish visa ineligibility."
    return {**meta, "checkedAt": dataset["checkedAt"], "status": "fresh" if dataset["occupationTableParsed"] else "partial",
            "freshness": "Published tables loaded", "latestRound": {**dataset["latestRound"], "minimumPoints": points,
            "scoreForOccupation": f"{points} points" if points is not None else None,
            "matchedOccupation": match["occupation"] if match else None, "occupationStatus": "published" if points is not None else "not_published" if dataset["occupationTableParsed"] else "unparsed"},
            "occupationMapping": mapping, "mappingNote": mapping_note, "occupationNote": note, "nextRound": dataset["nextRound"]}


def profile_checklist(profile):
    def supplied(value):
        return bool(clean(value)) and norm(value) not in {"unknown", "unsure", "n a", "na"}
    visa = norm(profile.get("visa"))
    temporary = bool(visa) and not any(word in visa for word in ("permanent", "citizen", "offshore", "no current australian visa"))
    experience = profile.get("experienceYears")
    rows = [
        ("occupation", "Target occupation", supplied(profile.get("occupation")) and norm(profile.get("occupation")) not in {"general professional", "skilled professional"}),
        ("classification", "Occupation code", bool(re.fullmatch(r"\d{6}", str(profile.get("anzsco") or "")) or re.fullmatch(r"\d{6}", str(profile.get("osca") or "")))),
        ("location", "Location", supplied(profile.get("location"))),
        ("education", "Qualification", supplied(profile.get("education")) and "detected from resume" not in str(profile.get("education", "")).lower()),
        ("experience", "Experience duration", isinstance(experience, (int, float)) and not isinstance(experience, bool) and experience >= 0),
        ("skills", "Skills", any(supplied(skill) and "add skills" not in skill.lower() for skill in profile.get("skills", []) if isinstance(skill, str))),
        ("visa", "Visa or residency status", supplied(profile.get("visa")) and "unsure" not in visa),
    ]
    if temporary:
        try:
            expiry = datetime.strptime(profile.get("visaExpiry", ""), "%Y-%m-%d").date()
            expiry_known = expiry >= datetime.now(timezone.utc).date()
        except (ValueError, TypeError):
            expiry_known = False
        rows.append(("visaExpiry", "Future visa expiry date", expiry_known))
    checks = [{"id": key, "label": label, "present": bool(present)} for key, label, present in rows]
    return {"completed": sum(row["present"] for row in checks), "total": len(checks), "checks": checks,
            "explanation": "Counts profile fields supplied. Information is self reported; this does not measure job prospects, visa eligibility or migration points."}
